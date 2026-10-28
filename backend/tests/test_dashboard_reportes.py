"""Panel de control, reportes (JSON, CSV, Excel, PDF) y auditoría."""

import csv
import io
import re
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from reportlab import rl_config

from app.service import reportes_pdf
from app.service.reportes_service import Columna, Reporte
from tests.ayuda import crear_categoria, crear_empresa, crear_producto, crear_sede, crear_usuario, entrada


async def base(cliente, db):
    admin = await crear_empresa(db, cliente)
    norte, sur = await crear_sede(admin, "Norte"), await crear_sede(admin, "Sur")
    cat = await crear_categoria(admin, "Papelería")
    p1 = await crear_producto(admin, cat["id"], "D-1", "Cuaderno argollado", costo=1000, precio_venta=2500)
    p2 = await crear_producto(admin, cat["id"], "D-2", "=cmd|' /C calc'!A0", precio_venta=500)  # texto malicioso y sin costo
    await entrada(admin, norte["id"], {p1["id"]: 40, p2["id"]: 10})
    await entrada(admin, sur["id"], {p1["id"]: 5})
    return admin, norte, sur, p1, p2


async def test_dashboard_sin_datos_no_inventa_cifras(cliente, db):
    admin = await crear_empresa(db, cliente)
    d = (await admin.get("/api/dashboard")).json()
    assert d["inventario"]["unidades_totales"] == 0 and d["inventario"]["valor_inventario"] == 0
    assert d["ventas"]["total"] == 0 and d["ventas"]["numero_ventas"] == 0 and d["ventas"]["top_productos"] == []
    assert len(d["ventas"]["serie_diaria"]) == 30 and all(x["total"] == 0 for x in d["ventas"]["serie_diaria"])
    assert d["alertas"]["total"] == 0 and d["prediccion"] is None


async def test_dashboard_cifras_reales(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    await admin.put("/api/inventario/niveles", {"sede_id": sur["id"], "producto_id": p1["id"], "stock_minimo": 10})
    v = (await admin.post("/api/ventas", {"sede_id": norte["id"], "items": [{"producto_id": p1["id"], "cantidad": 4}, {"producto_id": p2["id"], "cantidad": 2}]})).json()
    anulada = (await admin.post("/api/ventas", {"sede_id": norte["id"], "items": [{"producto_id": p1["id"], "cantidad": 1}]})).json()
    await admin.post(f"/api/ventas/{anulada['id']}/anular", {"motivo": "Error de digitación"})
    d = (await admin.get("/api/dashboard")).json()
    inv = d["inventario"]
    assert inv["unidades_totales"] == 40 - 4 + 5 + (10 - 2) - 0 and inv["unidades_totales"] == 49
    assert inv["valor_inventario"] == (36 + 5) * 1000  # solo productos CON costo
    assert inv["valor_completo"] is False and inv["productos_sin_costo"] == 1
    assert d["ventas"]["total"] == 4 * 2500 + 2 * 500 == v["total"]
    assert d["ventas"]["numero_ventas"] == 1 and d["ventas"]["anuladas"] == 1  # la anulada no suma
    assert d["ventas"]["top_productos"][0]["codigo"] == "D-1"
    assert sum(x["total"] for x in d["ventas"]["serie_diaria"]) == d["ventas"]["total"]
    assert d["alertas"]["total"] == 1 and d["alertas"]["principales"][0]["tipo"] == "inventario_bajo"
    assert {s["sede"]: s["unidades"] for s in d["por_sede"]} == {"Norte": 44, "Sur": 5}
    assert len(d["movimientos_recientes"]) == 7
    solo_sur = (await admin.get(f"/api/dashboard?sede_id={sur['id']}")).json()
    assert solo_sur["inventario"]["unidades_totales"] == 5 and solo_sur["ventas"]["total"] == 0


async def test_dashboard_respeta_permisos_y_sede(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    op = await crear_usuario(admin, "usuario_inventario", "op@andina.com", norte["id"])
    d = (await op.get("/api/dashboard")).json()
    assert d["inventario"]["unidades_totales"] == 50 and d["prediccion"] is None and d["recomendaciones_pendientes"] is None
    assert [s["sede"] for s in d["por_sede"]] == ["Norte"]
    analista = await crear_usuario(admin, "analista_compras", "an@andina.com")
    assert (await analista.get("/api/dashboard")).json()["recomendaciones_pendientes"] == 0
    assert (await op.get(f"/api/dashboard?sede_id={sur['id']}")).status_code == 403


async def test_reporte_inventario_json_csv_y_excel(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    r = (await admin.get("/api/reportes/inventario")).json()
    assert r["total"] == 3 and r["resumen"]["Unidades en existencia"] == 55
    assert r["resumen"]["Valor del inventario (solo productos con costo)"] == 45000
    csv_resp = await admin.get("/api/reportes/inventario?formato=csv")
    assert csv_resp.headers["content-disposition"].startswith("attachment; filename=")
    texto = csv_resp.content.decode("utf-8-sig")
    filas = list(csv.reader(io.StringIO(texto)))
    assert filas[0][:3] == ["Sede", "Código", "Producto"]
    # Inyección de fórmulas neutralizada
    assert any(c.startswith("'=") for f in filas for c in f) and not any(c.startswith("=") for f in filas for c in f)
    xlsx = await admin.get("/api/reportes/inventario?formato=xlsx")
    libro = load_workbook(io.BytesIO(xlsx.content))
    assert libro.sheetnames[-1] == "Resumen" and libro.worksheets[0].max_row == 4
    assert libro.worksheets[0]["F2"].value is not None and libro.worksheets[0]["A1"].value == "Sede"


async def test_reportes_exigen_permiso_y_quedan_auditados(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    op = await crear_usuario(admin, "usuario_inventario", "op@andina.com", norte["id"])
    assert (await op.get("/api/reportes/inventario")).status_code == 403
    assert (await op.get("/api/reportes/ventas")).status_code == 403
    assert (await op.get("/api/reportes/alertas")).status_code == 200
    tipos = {r["tipo"] for r in (await op.get("/api/reportes")).json()}
    assert tipos == {"movimientos", "alertas"}
    assert (await admin.get("/api/reportes/inexistente")).status_code == 404
    assert (await admin.get("/api/reportes/ventas?desde=2026-02-10&hasta=2026-01-01")).status_code == 400
    await admin.get("/api/reportes/movimientos?formato=csv")
    reg = (await admin.get("/api/auditoria?modulo=reportes")).json()
    assert reg["total"] == 1 and reg["items"][0]["accion"] == "exportar" and reg["items"][0]["detalle"]["formato"] == "csv"


async def test_reporte_de_ventas_excluye_anuladas_del_total(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    (await admin.post("/api/ventas", {"sede_id": norte["id"], "items": [{"producto_id": p1["id"], "cantidad": 2}]})).json()
    b = (await admin.post("/api/ventas", {"sede_id": norte["id"], "items": [{"producto_id": p1["id"], "cantidad": 3}]})).json()
    await admin.post(f"/api/ventas/{b['id']}/anular", {"motivo": "Cliente se arrepintió"})
    r = (await admin.get("/api/reportes/ventas")).json()
    assert r["total"] == 2 and r["resumen"]["Total vendido"] == 5000 and r["resumen"]["Unidades vendidas"] == 2


async def test_auditoria_visibilidad_por_permiso(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    enc = await crear_usuario(admin, "encargado_sede", "enc@andina.com", norte["id"])
    await entrada(enc, norte["id"], {p1["id"]: 1})
    todo = (await admin.get("/api/auditoria?tamano=100")).json()
    assert any(r["modulo"] == "auth" for r in todo["items"]) and todo["total"] > 5
    propio = (await enc.get("/api/auditoria?tamano=100")).json()
    assert propio["total"] >= 1 and all(r["modulo"] != "auth" and r["sede_id"] == norte["id"] for r in propio["items"])
    op = await crear_usuario(admin, "usuario_inventario", "op@andina.com", norte["id"])
    assert (await op.get("/api/auditoria")).status_code == 403
    filtrado = (await admin.get("/api/auditoria?q=enc@andina.com&modulo=inventario")).json()
    assert filtrado["total"] == 1
    opciones = (await admin.get("/api/auditoria/opciones")).json()
    assert "inventario" in opciones["modulos"] and "entrada" in opciones["acciones"]


# ------------------------------------------------------------------------------ PDF
@pytest.fixture
def pdf_legible(monkeypatch):
    """ReportLab comprime el contenido de cada página; sin comprimir se puede buscar el texto en los bytes."""
    monkeypatch.setattr(rl_config, "pageCompression", 0)
    monkeypatch.setattr(rl_config, "useA85", 0)


def _reporte(filas, columnas=None, resumen=None, truncado=False) -> Reporte:
    columnas = columnas or [
        Columna("sede", "Sede"), Columna("producto", "Producto"), Columna("cantidad", "Existencias", "entero"),
        Columna("valor", "Valor", "dinero"), Columna("error", "WAPE", "porcentaje"), Columna("dia", "Día", "fecha"),
    ]
    return Reporte("inventario", "Inventario por sede", columnas, filas, resumen or {}, len(filas), truncado)


def _pdf(reporte: Reporte) -> bytes:
    return reportes_pdf.a_pdf(reporte, empresa="Comercial Andina", usuario="Ana Pérez", filtros_texto="sede #1", generado=datetime(2026, 10, 26, 18, 50))


async def test_reporte_en_pdf_se_descarga_y_queda_auditado(cliente, db):
    admin, norte, sur, p1, p2 = await base(cliente, db)
    r = await admin.get("/api/reportes/inventario?formato=pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert re.fullmatch(r'attachment; filename="inventario-\d{4}-\d{2}-\d{2}\.pdf"', r.headers["content-disposition"])
    assert r.content.startswith(b"%PDF-") and r.content.rstrip().endswith(b"%%EOF")
    reg = (await admin.get("/api/auditoria?modulo=reportes")).json()
    assert reg["total"] == 1 and reg["items"][0]["detalle"]["formato"] == "pdf"
    # Los reportes exigen el mismo permiso que en pantalla, Excel y CSV
    op = await crear_usuario(admin, "usuario_inventario", "op@andina.com", norte["id"])
    assert (await op.get("/api/reportes/inventario?formato=pdf")).status_code == 403
    assert (await op.get("/api/reportes/alertas?formato=pdf")).status_code == 200
    assert (await admin.get("/api/reportes/inventario?formato=docx")).status_code == 422


def test_pdf_formatea_en_espanol_y_trae_el_encabezado(pdf_legible):
    filas = [{"sede": "Norte", "producto": "Cuaderno", "cantidad": 1234, "valor": 1500.5, "error": 0.254, "dia": date(2026, 10, 8)}]
    contenido = _pdf(_reporte(filas, resumen={"Unidades en existencia": 55, "Valor del inventario": 45000, "WAPE global": 0.31})).decode("latin-1")
    for texto in ("Comercial Andina", "Inventario por sede", "Ana P\\351rez", "26/10/2026 18:50", "1.234", "$ 1.500,50", "25,4 %", "08/10/2026",
                  "$ 45.000", "31,0 %", "P\\341gina 1 de 1"):
        assert texto in contenido, texto


def test_pdf_no_interpreta_marcado_ni_se_rompe_con_caracteres_raros(pdf_legible):
    filas = [{"sede": "Norte", "producto": "<b>&</b> piñón ₂ 🚲", "cantidad": 1, "valor": None, "error": None, "dia": None}]
    contenido = _pdf(_reporte(filas)).decode("latin-1")
    assert contenido.startswith("%PDF-")
    assert "(<) Tj (b) Tj (>) Tj" in contenido  # las etiquetas se dibujan como texto, no se interpretan como negrita
    assert "pi\\361\\363n ? ?" in contenido  # ñ y ó se conservan; ₂ y 🚲 no existen en la fuente y salen como «?»
    assert contenido.count("\\227") >= 3  # los vacíos salen como «—» (en WinAnsi, 227 en octal)


def test_pdf_repite_el_encabezado_y_numera_las_paginas(pdf_legible):
    filas = [{"sede": "Norte", "producto": f"Producto {i}", "cantidad": i, "valor": 100.0 * i, "error": 0.1, "dia": date(2026, 10, 8)} for i in range(160)]
    contenido = _pdf(_reporte(filas)).decode("latin-1")
    paginas = len(re.findall(rb"/Type /Page\b".decode(), contenido))
    assert paginas >= 3
    assert f"P\\341gina 1 de {paginas}" in contenido and f"P\\341gina {paginas} de {paginas}" in contenido
    assert contenido.count("Existencias") == paginas  # el encabezado de la tabla se repite en cada página


def test_pdf_avisa_cuando_se_recorta_y_cuando_no_hay_datos(pdf_legible):
    fila = {"sede": "Norte", "producto": "A", "cantidad": 1, "valor": 1.0, "error": 0.1, "dia": date(2026, 10, 8)}
    recortado = _pdf(_reporte([fila] * 5, truncado=True)).decode("latin-1")
    assert "solo las primeras 5 filas" in recortado and "Excel o CSV" in recortado
    assert "el resumen de arriba suma solo esas filas" in recortado
    vacio = _pdf(_reporte([])).decode("latin-1")
    assert "No hay datos para los filtros elegidos" in vacio


def test_pdf_horizontal_cuando_hay_muchas_columnas():
    muchas = [Columna(f"c{i}", f"Columna {i}") for i in range(9)]
    pocas = [Columna(f"c{i}", f"Columna {i}") for i in range(3)]
    fila = {f"c{i}": "x" for i in range(9)}
    assert b"/MediaBox [ 0 0 841.8" in _pdf(_reporte([fila], muchas))  # A4 horizontal
    assert b"/MediaBox [ 0 0 595.2" in _pdf(_reporte([fila], pocas))  # A4 vertical


@pytest.mark.parametrize("valor,tipo,esperado", [
    (1234567, "entero", "1.234.567"), (1500.5, "decimal", "1.500,50"), (0.254, "porcentaje", "25,4 %"),
    (2500, "dinero", "$ 2.500"), (None, "entero", "—"), (date(2026, 1, 5), "fecha", "05/01/2026"),
    (datetime(2026, 1, 5, 9, 7), "fechahora", "05/01/2026 09:07"), ("texto", "texto", "texto"),
])
def test_pdf_formato_de_valores(valor, tipo, esperado):
    assert reportes_pdf.formatear(valor, tipo) == esperado


@pytest.mark.postgres
@pytest.mark.parametrize("formato", ["csv", "xlsx", "pdf"])
async def test_exportar_con_una_sede_inexistente_no_falla_y_queda_auditado_sin_sede(cliente, db, formato):
    """Lo encontró el fuzzing: la auditoría apuntaba a una sede que no existe y la clave foránea producía un 500."""
    from sqlalchemy import select

    from app.model.auditoria import Auditoria

    admin = await crear_empresa(db, cliente)
    r = await admin.get("/api/reportes/traslados", params={"formato": formato, "sede_id": 197519})
    assert r.status_code == 200, r.text
    registro = (await db.scalars(select(Auditoria).where(Auditoria.modulo == "reportes").order_by(Auditoria.id.desc()))).first()
    assert registro is not None and registro.accion == "exportar" and registro.sede_id is None
    assert "197519" in registro.detalle["filtros"]
