"""Reportes de la plataforma (RF-034, RF-035): vista previa en pantalla y exportación a CSV y Excel.

Cada reporte se arma con consultas reales, respeta la empresa y la sede del usuario y exige su permiso
(``reportes.<tipo>``). Las exportaciones quedan en la auditoría. Los CSV neutralizan fórmulas de hoja de
cálculo (``=``, ``+``, ``-``, ``@``) para evitar la inyección de fórmulas al abrirlos en Excel."""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.acceso import Contexto
from app.core.exceptions import NoEncontrado, ReglaNegocio
from app.core.niveles import clasificar_stock
from app.core.tiempo import ZONA, ahora_utc, asegurar_utc, fin_de_dia_utc, inicio_de_dia_utc
from app.models.catalogo import Categoria, Producto
from app.models.empresa import Empresa
from app.models.inventario import Alerta, InventarioSede, Movimiento
from app.models.operaciones import Traslado, Venta, VentaItem
from app.models.prediccion import EjecucionPrediccion, Prediccion, Recomendacion
from app.models.sede import Sede
from app.models.usuario import Usuario
from app.services import auditoria_service, inventario_service
from app.services.alcance import resolver_filtro_sede

LIMITE_VISTA_PREVIA = 200
LIMITE_EXPORTACION = 50_000


@dataclass
class Columna:
    clave: str
    titulo: str
    tipo: str = "texto"  # texto | entero | dinero | decimal | fecha | fechahora | porcentaje


@dataclass
class Reporte:
    tipo: str
    titulo: str
    columnas: list[Columna]
    filas: list[dict[str, Any]]
    resumen: dict[str, Any] = field(default_factory=dict)
    total: int = 0
    truncado: bool = False


@dataclass
class Filtros:
    sede_id: int | None = None
    desde: date | None = None
    hasta: date | None = None
    categoria_id: int | None = None
    estado: str | None = None
    tipo: str | None = None


TIPOS = {
    "inventario": ("reportes.inventario", "Inventario por sede"),
    "ventas": ("reportes.ventas", "Ventas"),
    "movimientos": ("reportes.movimientos", "Movimientos de inventario"),
    "traslados": ("reportes.traslados", "Traslados entre sedes"),
    "alertas": ("reportes.alertas", "Alertas de inventario"),
    "predicciones": ("reportes.predicciones", "Predicción de demanda"),
    "recomendaciones": ("reportes.recomendaciones", "Recomendaciones de reabastecimiento y traslado"),
}


def _local(momento: datetime | None) -> datetime | None:
    if momento is None:
        return None
    return asegurar_utc(momento).astimezone(ZONA).replace(tzinfo=None)


def _num(valor) -> float | None:
    return None if valor is None else float(valor)


# ------------------------------------------------------------------------------ constructores


async def _inventario(db, ctx, f: Filtros, limite) -> Reporte:
    sede_ids = resolver_filtro_sede(ctx, f.sede_id)
    consulta = (
        select(InventarioSede, Producto, Categoria.nombre, Sede.nombre)
        .join(Producto, Producto.id == InventarioSede.producto_id)
        .join(Categoria, Categoria.id == Producto.categoria_id)
        .join(Sede, Sede.id == InventarioSede.sede_id)
        .where(InventarioSede.empresa_id == ctx.empresa_id)
    )
    if sede_ids is not None:
        consulta = consulta.where(InventarioSede.sede_id.in_(sede_ids))
    if f.categoria_id:
        consulta = consulta.where(Producto.categoria_id == f.categoria_id)
    filas, valor, unidades, sin_costo = [], Decimal(0), 0, 0
    for inv, prod, cat, sede in (await db.execute(consulta.order_by(Sede.nombre, Producto.nombre))).all():
        estado = clasificar_stock(inv.cantidad, inv.stock_minimo, inv.stock_maximo)
        if f.estado and estado != f.estado:
            continue
        v = prod.costo * inv.cantidad if prod.costo is not None else None
        valor += v or 0
        unidades += inv.cantidad
        sin_costo += 1 if (prod.costo is None and inv.cantidad > 0) else 0
        filas.append({
            "sede": sede, "codigo": prod.codigo, "producto": prod.nombre, "categoria": cat, "unidad": prod.unidad_medida,
            "cantidad": inv.cantidad, "stock_minimo": inv.stock_minimo, "stock_maximo": inv.stock_maximo, "estado": estado,
            "costo": _num(prod.costo), "valor": _num(v),
        })
    cols = [Columna("sede", "Sede"), Columna("codigo", "Código"), Columna("producto", "Producto"), Columna("categoria", "Categoría"),
            Columna("unidad", "Unidad"), Columna("cantidad", "Existencias", "entero"), Columna("stock_minimo", "Stock mínimo", "entero"),
            Columna("stock_maximo", "Stock máximo", "entero"), Columna("estado", "Estado"), Columna("costo", "Costo unitario", "dinero"),
            Columna("valor", "Valor (costo × existencias)", "dinero")]
    return Reporte("inventario", TIPOS["inventario"][1], cols, filas, {
        "Unidades en existencia": unidades, "Valor del inventario (solo productos con costo)": float(valor),
        "Registros sin costo (no valorados)": sin_costo,
    }, total=len(filas))


async def _ventas(db, ctx, f: Filtros, limite) -> Reporte:
    sede_ids = resolver_filtro_sede(ctx, f.sede_id)
    consulta = (
        select(Venta, VentaItem, Producto, Sede.nombre, Usuario.nombre)
        .join(VentaItem, VentaItem.venta_id == Venta.id).join(Producto, Producto.id == VentaItem.producto_id)
        .join(Sede, Sede.id == Venta.sede_id).join(Usuario, Usuario.id == Venta.usuario_id)
        .where(Venta.empresa_id == ctx.empresa_id)
    )
    if sede_ids is not None:
        consulta = consulta.where(Venta.sede_id.in_(sede_ids))
    if f.desde:
        consulta = consulta.where(Venta.fecha_dia >= f.desde)
    if f.hasta:
        consulta = consulta.where(Venta.fecha_dia <= f.hasta)
    if f.estado:
        consulta = consulta.where(Venta.estado == f.estado)
    consulta = consulta.order_by(Venta.fecha.desc(), Venta.id.desc(), VentaItem.id).limit(limite + 1)
    resultado = (await db.execute(consulta)).all()
    truncado = len(resultado) > limite
    filas, total_vendido, unidades, ventas_ids = [], Decimal(0), 0, set()
    for v, it, prod, sede, vendedor in resultado[:limite]:
        filas.append({
            "numero": v.numero, "fecha": _local(v.fecha), "sede": sede, "vendedor": vendedor, "estado": v.estado,
            "codigo": prod.codigo, "producto": prod.nombre, "cantidad": it.cantidad, "precio": _num(it.precio_unitario),
            "subtotal": _num(it.subtotal),
        })
        if v.estado == "registrada":
            total_vendido += it.subtotal
            unidades += it.cantidad
            ventas_ids.add(v.id)
    cols = [Columna("numero", "N.º de venta", "entero"), Columna("fecha", "Fecha", "fechahora"), Columna("sede", "Sede"),
            Columna("vendedor", "Registrada por"), Columna("estado", "Estado"), Columna("codigo", "Código"), Columna("producto", "Producto"),
            Columna("cantidad", "Cantidad", "entero"), Columna("precio", "Precio unitario", "dinero"), Columna("subtotal", "Subtotal", "dinero")]
    return Reporte("ventas", TIPOS["ventas"][1], cols, filas, {
        "Ventas registradas (sin anuladas)": len(ventas_ids), "Unidades vendidas": unidades, "Total vendido": float(total_vendido),
    }, total=len(filas), truncado=truncado)


async def _movimientos(db, ctx, f: Filtros, limite) -> Reporte:
    condiciones = inventario_service.condiciones_movimientos(
        ctx, sede_id=f.sede_id, producto_id=None, tipo=f.tipo, desde=f.desde, hasta=f.hasta, usuario_id=None
    )
    movs = (await db.scalars(select(Movimiento).where(*condiciones).order_by(Movimiento.fecha.desc(), Movimiento.id.desc()).limit(limite + 1))).all()
    truncado = len(movs) > limite
    salidas = await inventario_service.movimientos_a_salidas(db, list(movs[:limite]))
    filas = [{
        "fecha": _local(m.fecha), "sede": m.sede.nombre, "codigo": m.producto.codigo, "producto": m.producto.nombre, "tipo": m.tipo,
        "motivo": m.motivo, "delta": m.delta, "saldo_anterior": m.saldo_anterior, "saldo_posterior": m.saldo_posterior,
        "usuario": m.usuario_nombre, "referencia": (f"Venta #{m.venta_numero}" if m.venta_numero else f"Traslado #{m.traslado_numero}" if m.traslado_numero else None),
        "observacion": m.observacion,
    } for m in salidas]
    cols = [Columna("fecha", "Fecha", "fechahora"), Columna("sede", "Sede"), Columna("codigo", "Código"), Columna("producto", "Producto"),
            Columna("tipo", "Tipo"), Columna("motivo", "Motivo"), Columna("delta", "Variación", "entero"), Columna("saldo_anterior", "Saldo anterior", "entero"),
            Columna("saldo_posterior", "Saldo posterior", "entero"), Columna("usuario", "Registrado por"), Columna("referencia", "Referencia"),
            Columna("observacion", "Observación")]
    return Reporte("movimientos", TIPOS["movimientos"][1], cols, filas, {
        "Entradas (unidades)": sum(x["delta"] for x in filas if x["delta"] > 0),
        "Salidas (unidades)": -sum(x["delta"] for x in filas if x["delta"] < 0),
    }, total=len(filas), truncado=truncado)


async def _traslados(db, ctx, f: Filtros, limite) -> Reporte:
    autorizadas = ctx.sedes_autorizadas()
    consulta = select(Traslado).where(Traslado.empresa_id == ctx.empresa_id)
    if autorizadas is not None:
        consulta = consulta.where(Traslado.sede_origen_id.in_(autorizadas) | Traslado.sede_destino_id.in_(autorizadas))
    if f.sede_id:
        consulta = consulta.where((Traslado.sede_origen_id == f.sede_id) | (Traslado.sede_destino_id == f.sede_id))
    if f.estado:
        consulta = consulta.where(Traslado.estado == f.estado)
    if f.desde:
        consulta = consulta.where(Traslado.solicitado_en >= inicio_de_dia_utc(f.desde))
    if f.hasta:
        consulta = consulta.where(Traslado.solicitado_en < fin_de_dia_utc(f.hasta))
    traslados = (await db.scalars(consulta.order_by(Traslado.solicitado_en.desc()).limit(limite + 1))).all()
    truncado = len(traslados) > limite
    traslados = traslados[:limite]
    sedes = {s.id: s.nombre for s in (await db.scalars(select(Sede).where(Sede.empresa_id == ctx.empresa_id))).all()}
    usuarios = {u.id: u.nombre for u in (await db.execute(select(Usuario.id, Usuario.nombre).where(Usuario.empresa_id == ctx.empresa_id))).all()}
    productos = {p.id: p for p in (await db.scalars(select(Producto).where(Producto.empresa_id == ctx.empresa_id))).all()}
    filas = []
    for t in traslados:
        filas.append({
            "numero": t.numero, "estado": t.estado, "origen": sedes.get(t.sede_origen_id), "destino": sedes.get(t.sede_destino_id),
            "productos": "; ".join(f"{productos[i.producto_id].codigo} × {i.cantidad}" for i in t.items), "unidades": sum(i.cantidad for i in t.items),
            "solicitado": _local(t.solicitado_en), "solicitado_por": usuarios.get(t.solicitado_por), "aprobado": _local(t.aprobado_en),
            "confirmado": _local(t.confirmado_en), "motivo_cierre": t.motivo_cierre,
        })
    cols = [Columna("numero", "N.º", "entero"), Columna("estado", "Estado"), Columna("origen", "Sede de origen"), Columna("destino", "Sede de destino"),
            Columna("productos", "Productos"), Columna("unidades", "Unidades", "entero"), Columna("solicitado", "Solicitado", "fechahora"),
            Columna("solicitado_por", "Solicitado por"), Columna("aprobado", "Aprobado", "fechahora"), Columna("confirmado", "Confirmado", "fechahora"),
            Columna("motivo_cierre", "Motivo de cierre")]
    return Reporte("traslados", TIPOS["traslados"][1], cols, filas, {
        "Traslados confirmados": sum(1 for x in filas if x["estado"] == "confirmado"),
        "Pendientes": sum(1 for x in filas if x["estado"] in ("solicitado", "aprobado")),
    }, total=len(filas), truncado=truncado)


async def _alertas(db, ctx, f: Filtros, limite) -> Reporte:
    sede_ids = resolver_filtro_sede(ctx, f.sede_id)
    consulta = (
        select(Alerta, Producto, Sede.nombre).join(Producto, Producto.id == Alerta.producto_id).join(Sede, Sede.id == Alerta.sede_id)
        .where(Alerta.empresa_id == ctx.empresa_id)
    )
    if sede_ids is not None:
        consulta = consulta.where(Alerta.sede_id.in_(sede_ids))
    if f.estado:
        consulta = consulta.where(Alerta.estado == f.estado)
    if f.tipo:
        consulta = consulta.where(Alerta.tipo == f.tipo)
    if f.desde:
        consulta = consulta.where(Alerta.abierta_en >= inicio_de_dia_utc(f.desde))
    if f.hasta:
        consulta = consulta.where(Alerta.abierta_en < fin_de_dia_utc(f.hasta))
    res = (await db.execute(consulta.order_by(Alerta.abierta_en.desc()).limit(limite + 1))).all()
    truncado = len(res) > limite
    filas = [{
        "tipo": a.tipo, "prioridad": a.prioridad, "estado": a.estado, "sede": sede, "codigo": p.codigo, "producto": p.nombre,
        "mensaje": a.mensaje, "abierta": _local(a.abierta_en), "resuelta": _local(a.resuelta_en),
    } for a, p, sede in res[:limite]]
    cols = [Columna("tipo", "Tipo"), Columna("prioridad", "Prioridad"), Columna("estado", "Estado"), Columna("sede", "Sede"),
            Columna("codigo", "Código"), Columna("producto", "Producto"), Columna("mensaje", "Mensaje"),
            Columna("abierta", "Abierta", "fechahora"), Columna("resuelta", "Resuelta", "fechahora")]
    return Reporte("alertas", TIPOS["alertas"][1], cols, filas, {"Alertas activas": sum(1 for x in filas if x["estado"] == "activa")},
                   total=len(filas), truncado=truncado)


async def _predicciones(db, ctx, f: Filtros, limite) -> Reporte:
    sede_ids = resolver_filtro_sede(ctx, f.sede_id)
    ej = await db.scalar(
        select(EjecucionPrediccion).where(EjecucionPrediccion.empresa_id == ctx.empresa_id, EjecucionPrediccion.estado == "completada")
        .order_by(EjecucionPrediccion.finalizada_en.desc()).limit(1)
    )
    if ej is None:
        return Reporte("predicciones", TIPOS["predicciones"][1], [], [], {"Aviso": "Aún no hay predicciones generadas"}, 0)
    consulta = (
        select(Prediccion, Producto, Sede.nombre).join(Producto, Producto.id == Prediccion.producto_id).join(Sede, Sede.id == Prediccion.sede_id)
        .where(Prediccion.ejecucion_id == ej.id)
    )
    if sede_ids is not None:
        consulta = consulta.where(Prediccion.sede_id.in_(sede_ids))
    res = (await db.execute(consulta.order_by(Sede.nombre, Producto.nombre).limit(limite + 1))).all()
    truncado = len(res) > limite
    filas = []
    for p, prod, sede in res[:limite]:
        m = p.metricas or {}
        filas.append({
            "sede": sede, "codigo": prod.codigo, "producto": prod.nombre, "desde": p.fecha_inicio, "hasta": p.fecha_fin,
            "demanda": _num(p.demanda_estimada), "inferior": _num(p.limite_inferior), "superior": _num(p.limite_superior),
            "modelo": p.modelo, "confianza": p.confianza, "wape": m.get("wape"), "mae": m.get("mae"), "wape_base": m.get("wape_linea_base"),
            "advertencias": "; ".join(p.advertencias or []),
        })
    cols = [Columna("sede", "Sede"), Columna("codigo", "Código"), Columna("producto", "Producto"), Columna("desde", "Desde", "fecha"),
            Columna("hasta", "Hasta", "fecha"), Columna("demanda", "Demanda estimada (unidades)", "decimal"),
            Columna("inferior", "Escenario bajo", "decimal"), Columna("superior", "Escenario alto", "decimal"), Columna("modelo", "Modelo"),
            Columna("confianza", "Confianza"), Columna("wape", "WAPE del modelo", "porcentaje"), Columna("mae", "MAE (unid./día)", "decimal"),
            Columna("wape_base", "WAPE línea base", "porcentaje"), Columna("advertencias", "Advertencias")]
    r = ej.resumen or {}
    return Reporte("predicciones", TIPOS["predicciones"][1], cols, filas, {
        "Ejecución": f"#{ej.id} · datos hasta {ej.fecha_corte}", "Series con predicción": r.get("series_con_prediccion"),
        "Series omitidas por datos insuficientes": r.get("series_omitidas"), "WAPE global": r.get("wape_global"),
        "Nota": "Son estimaciones, no ventas reales.",
    }, total=len(filas), truncado=truncado)


async def _recomendaciones(db, ctx, f: Filtros, limite) -> Reporte:
    sede_ids = resolver_filtro_sede(ctx, f.sede_id)
    consulta = (
        select(Recomendacion, Producto, Sede.nombre).join(Producto, Producto.id == Recomendacion.producto_id)
        .join(Sede, Sede.id == Recomendacion.sede_id).where(Recomendacion.empresa_id == ctx.empresa_id)
    )
    if sede_ids is not None:
        consulta = consulta.where(Recomendacion.sede_id.in_(sede_ids))
    if f.estado:
        consulta = consulta.where(Recomendacion.estado == f.estado)
    if f.tipo:
        consulta = consulta.where(Recomendacion.tipo == f.tipo)
    res = (await db.execute(consulta.order_by(Recomendacion.generada_en.desc(), Recomendacion.id.desc()).limit(limite + 1))).all()
    truncado = len(res) > limite
    origenes = {s.id: s.nombre for s in (await db.scalars(select(Sede).where(Sede.empresa_id == ctx.empresa_id))).all()}
    filas = [{
        "tipo": r.tipo, "prioridad": r.prioridad, "estado": r.estado, "sede": sede, "origen": origenes.get(r.sede_origen_id),
        "codigo": p.codigo, "producto": p.nombre, "cantidad": r.cantidad_sugerida, "resumen": r.resumen,
        "generada": _local(r.generada_en), "revisada": _local(r.revisada_en), "comentario": r.comentario_revision,
    } for r, p, sede in res[:limite]]
    cols = [Columna("tipo", "Tipo"), Columna("prioridad", "Prioridad"), Columna("estado", "Estado"), Columna("sede", "Sede que lo necesita"),
            Columna("origen", "Sede de origen (traslados)"), Columna("codigo", "Código"), Columna("producto", "Producto"),
            Columna("cantidad", "Cantidad sugerida", "entero"), Columna("resumen", "Explicación"), Columna("generada", "Generada", "fechahora"),
            Columna("revisada", "Revisada", "fechahora"), Columna("comentario", "Comentario de la revisión")]
    return Reporte("recomendaciones", TIPOS["recomendaciones"][1], cols, filas, {
        "Pendientes de revisión": sum(1 for x in filas if x["estado"] == "pendiente")}, total=len(filas), truncado=truncado)


_CONSTRUCTORES = {
    "inventario": _inventario, "ventas": _ventas, "movimientos": _movimientos, "traslados": _traslados,
    "alertas": _alertas, "predicciones": _predicciones, "recomendaciones": _recomendaciones,
}


async def generar(db: AsyncSession, ctx: Contexto, tipo: str, filtros: Filtros, *, limite: int) -> Reporte:
    if tipo not in TIPOS:
        raise NoEncontrado("Ese reporte no existe")
    ctx.exigir(TIPOS[tipo][0])
    if filtros.desde and filtros.hasta and filtros.desde > filtros.hasta:
        raise ReglaNegocio("La fecha «desde» no puede ser posterior a «hasta»", codigo="rango_invalido")
    return await _CONSTRUCTORES[tipo](db, ctx, filtros, limite)


# ------------------------------------------------------------------------------ exportación


def _neutralizar(valor: Any) -> Any:
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


def _formato_csv(valor: Any, tipo: str) -> Any:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    if isinstance(valor, date):
        return valor.isoformat()
    if tipo == "porcentaje" and isinstance(valor, (int, float)):
        return f"{valor * 100:.1f}%"
    return _neutralizar(valor)


def a_csv(reporte: Reporte) -> bytes:
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=",", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    escritor.writerow([c.titulo for c in reporte.columnas])
    for fila in reporte.filas:
        escritor.writerow([_formato_csv(fila.get(c.clave), c.tipo) for c in reporte.columnas])
    return ("﻿" + salida.getvalue()).encode("utf-8")  # BOM: Excel reconoce las tildes y la ñ


_FORMATOS_XLSX = {
    "entero": "#,##0", "dinero": '"$"#,##0.00', "decimal": "#,##0.00", "porcentaje": "0.0%",
    "fecha": "yyyy-mm-dd", "fechahora": "yyyy-mm-dd hh:mm",
}


def a_xlsx(reporte: Reporte, *, empresa: str, usuario: str, filtros_texto: str) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte.titulo[:31]
    relleno = PatternFill("solid", start_color="1F2A44")
    borde = Border(bottom=Side(style="thin", color="D0D5DD"))
    for j, col in enumerate(reporte.columnas, start=1):
        celda = hoja.cell(row=1, column=j, value=col.titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for i, fila in enumerate(reporte.filas, start=2):
        for j, col in enumerate(reporte.columnas, start=1):
            valor = fila.get(col.clave)
            celda = hoja.cell(row=i, column=j, value=_neutralizar(valor) if isinstance(valor, str) else valor)
            if col.tipo in _FORMATOS_XLSX and valor is not None:
                celda.number_format = _FORMATOS_XLSX[col.tipo]
            celda.border = borde
    for j, col in enumerate(reporte.columnas, start=1):
        mayor = max([len(col.titulo)] + [len(str(f.get(col.clave) or "")) for f in reporte.filas[:300]])
        hoja.column_dimensions[get_column_letter(j)].width = min(max(10, mayor + 2), 60)
    hoja.freeze_panes = "A2"
    if reporte.columnas and reporte.filas:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(reporte.columnas))}{len(reporte.filas) + 1}"

    info = libro.create_sheet("Resumen")
    info["A1"], info["A1"].font = reporte.titulo, Font(bold=True, size=14)
    meta = [("Empresa", empresa), ("Generado por", usuario), ("Generado el", _local(ahora_utc())), ("Filtros", filtros_texto or "Sin filtros"),
            ("Filas", reporte.total)] + list(reporte.resumen.items())
    if reporte.truncado:
        meta.append(("Aviso", f"El reporte se recortó a {reporte.total} filas; acota las fechas para ver el resto"))
    for i, (k, v) in enumerate(meta, start=3):
        info.cell(row=i, column=1, value=k).font = Font(bold=True)
        info.cell(row=i, column=2, value=v)
        if isinstance(v, datetime):
            info.cell(row=i, column=2).number_format = "yyyy-mm-dd hh:mm"
    info.column_dimensions["A"].width = 46
    info.column_dimensions["B"].width = 60
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def descripcion_filtros(f: Filtros) -> str:
    partes = []
    if f.sede_id:
        partes.append(f"sede #{f.sede_id}")
    if f.desde:
        partes.append(f"desde {f.desde}")
    if f.hasta:
        partes.append(f"hasta {f.hasta}")
    if f.categoria_id:
        partes.append(f"categoría #{f.categoria_id}")
    if f.estado:
        partes.append(f"estado {f.estado}")
    if f.tipo:
        partes.append(f"tipo {f.tipo}")
    return ", ".join(partes)


async def auditar_exportacion(db: AsyncSession, ctx: Contexto, tipo: str, formato: str, filtros: Filtros, filas: int) -> None:
    await auditoria_service.registrar(
        db, modulo="reportes", accion="exportar", ctx=ctx, entidad="reporte", entidad_id=tipo,
        sede_id=ctx.sede_id if ctx.limitado_a_sede else filtros.sede_id,
        detalle={"formato": formato, "filas": filas, "filtros": descripcion_filtros(filtros)},
    )
    await db.commit()


async def nombre_empresa(db: AsyncSession, ctx: Contexto) -> str:
    return (await db.get(Empresa, ctx.empresa_id)).nombre
